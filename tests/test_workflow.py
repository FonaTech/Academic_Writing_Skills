"""Behavioral regressions for observed failures, provenance and declarative builds.
Run: python3 -m unittest discover -s tests -v
Synthetic inputs are test fixtures, not research findings.
"""
from __future__ import annotations
import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
LIT = ROOT/'skills/sci-literature-evidence/scripts'
KIT = ROOT/'skills/sci-manuscript-build/assets/manuscript_kit'
sys.path.insert(0,str(LIT))
from evidence_tools import verify_quote, numeric_tokens, quantities, audit_bindings, verify_record, sha256

def write(path,data):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

def run(path,*args,cwd=None):
    return subprocess.run([sys.executable,str(path),*map(str,args)],cwd=cwd,capture_output=True,text=True,timeout=60)

class EvidenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        (self.root/'source.txt').write_text('=== PAGE 1 ===\nMeasured efficiency was 1.5 TOPS/W.\n=== PAGE 2 ===\nThe next sentence continues here.',encoding='utf-8')
        self.rec = dict(source_file=str(self.root/'source.txt'),page=1,quote='Measured efficiency was 1.5 TOPS/W.',value='1.5')

    def tearDown(self):
        self.temp.cleanup()

    def test_complete_quote_and_decimal(self):
        self.assertEqual(verify_quote(self.rec,self.root),[])
        wrong = dict(self.rec,quote='Measured efficiency was 15 TOPS/W.',value='15')
        self.assertTrue(verify_quote(wrong,self.root))

    def test_complete_tail(self):
        prefix = 'This synthetic source has a long explanatory opening. '*5
        (self.root/'source.txt').write_text('=== PAGE 1 ===\n'+prefix+'The value was 1.5 V.',encoding='utf-8')
        rec = dict(self.rec,quote=prefix+'The value was 900 V.',value='900')
        self.assertTrue(verify_quote(rec,self.root))

    def test_page_and_missing_fields(self):
        self.assertTrue(verify_quote(dict(self.rec,page=2),self.root))
        self.assertTrue(verify_quote(dict(self.rec,page=9),self.root))
        self.assertTrue(verify_quote(dict(self.rec,quote=''),self.root))
        self.assertTrue(verify_quote(dict(self.rec,page=None),self.root))

    def test_cross_page_quote(self):
        rec = dict(self.rec,pages=[1,2],quote='1.5 TOPS/W. The next sentence continues here.')
        self.assertEqual(verify_quote(rec,self.root),[])

    def test_numeric_token_not_substring(self):
        self.assertTrue(verify_quote(dict(self.rec,value='5'),self.root))
        self.assertNotEqual(numeric_tokens('-2 V'),numeric_tokens('2 V'))
        self.assertNotEqual(numeric_tokens('1e-7 V'),numeric_tokens('1e7 V'))
        self.assertEqual(numeric_tokens('10⁻⁷ V'),numeric_tokens('1e-7 V'))
        self.assertEqual(quantities('2 mg/L'),[('2','mg/L')])
        self.assertEqual(quantities('59–405 mV'),[('59','mV'),('405','mV')])

    def test_empty_record_set_fails(self):
        file = self.root/'empty.json';write(file,{})
        self.assertNotEqual(run(LIT/'verify_evidence_quotes.py',file,'--text-dir',self.root).returncode,0)

    def test_wrong_metric_same_number_is_not_support(self):
        rec = dict(value='62.11',unit='TOPS/W',metric='efficiency',entity='B',result_method='measured',source_kind='raw_data')
        claim = dict(evidence_key='b',value='62.11',unit='J',metric='energy',entity='A',result_method='measured',text={'en':'Device A consumed 62.11 J.'})
        block = dict(claims=[claim])
        errors,_,_ = audit_bindings([('p1',block,{'en':claim['text']['en']})],{'b':rec},release=True)
        self.assertGreaterEqual(len(errors),3)

    def test_no_binding_fails_even_when_number_exists(self):
        errors,_,_ = audit_bindings([('p1',{},dict(en='The voltage was 2 V.'))],{'other':dict(value='2')},release=True)
        self.assertTrue(errors)

    def test_provenance_hash_change(self):
        file = self.root/'raw.json';write(file,{'value':2})
        rec = dict(source_kind='raw_data',provenance=dict(artifacts=[dict(path='raw.json',sha256=sha256(file))],locator='value'))
        self.assertEqual(verify_record(rec,self.root,self.root),[])
        write(file,{'value':3})
        self.assertTrue(verify_record(rec,self.root,self.root))

    def test_table_header_unit_inheritance_and_wrong_header(self):
        rec = dict(value='21',unit='°C',metric='temperature',entity='A',result_method='synthetic',source_kind='raw_data')
        claim = dict(rec,evidence_key='a',text={'en':'21'})
        block = dict(table_cell=True,unit_context={'en':'Mean temperature (°C)'},claims=[claim])
        errors,_,_=audit_bindings([('tab:t:0:1',block,{'en':'21'})],{'a':rec},release=True)
        self.assertEqual(errors,[])
        block['unit_context']['en']='Mean temperature (K)'
        self.assertTrue(audit_bindings([('cell',block,{'en':'21'})],{'a':rec},release=True)[0])
        block['unit_context']['en']='Mean temperature (°C)'
        claim['text']['en']='21 mV'
        self.assertTrue(audit_bindings([('cell',block,{'en':'21 mV'})],{'a':rec},release=True)[0])

    def test_unbound_numeric_table_cell(self):
        block=dict(table_cell=True,unit_context={'en':'Number of records'})
        self.assertTrue(audit_bindings([('cell',block,{'en':'2'})],{},release=True)[0])

    def test_declared_units_across_fields(self):
        # Artificial values test unit mechanics only, never field-specific inference.
        for field,unit in [('clinical','mmHg'),('education','points'),('economics','USD'),('chemistry','mg/L'),('biology','bpm'),('computing','tokens/s')]:
            with self.subTest(field=field):
                text=f'Synthetic fixture result: 2 {unit}.'
                rec=dict(source_kind='raw_data',value='2',unit=unit,metric='fixture metric',entity=field,result_method='synthetic')
                claim=dict(rec,evidence_key='v',text={'en':text})
                self.assertEqual(audit_bindings([('p',dict(claims=[claim]),{'en':text})],{'v':rec},release=True)[0],[])

class ProjectTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)/'paper'
        result = run(ROOT/'skills/sci-manuscript-build/scripts/new_project.py',self.root,'--slug','Fixture','--type','article','--languages','en','zh')
        self.assertEqual(result.returncode,0,result.stderr)
        write(self.root/'content.json',dict(schema_version=2,chapter=None,blocks=[dict(id='p1',type='para',en='The synthetic fixture voltage was 2 V.',zh='人工合成测试数据的电压为 2 V。')]))

    def tearDown(self):
        self.temp.cleanup()

    def build(self):
        result = run(self.root/'build_manuscript.py',cwd=self.root)
        self.assertEqual(result.returncode,0,result.stderr)

    def check(self,*args):
        return run(self.root/'check_manuscript.py',*args,cwd=self.root)

    def test_small_bilingual_value_difference(self):
        content = json.loads((self.root/'content.json').read_text())
        content['blocks'][0]['zh']='人工合成测试数据的电压为 3 V。'
        write(self.root/'content.json',content);self.build()
        result = self.check()
        self.assertNotEqual(result.returncode,0)
        self.assertIn('differ',result.stdout)

    def test_bilingual_unit_difference(self):
        content = json.loads((self.root/'content.json').read_text())
        content['blocks'][0]['zh']='人工合成测试数据的电压为 2 mV。'
        write(self.root/'content.json',content);self.build()
        self.assertNotEqual(self.check().returncode,0)

    def test_chinese_primary_does_not_overwrite_english(self):
        cfg = json.loads((self.root/'manuscript.json').read_text());cfg['output']['languages']=['zh','en'];write(self.root/'manuscript.json',cfg)
        self.build()
        from docx import Document
        en=' '.join(p.text for p in Document(self.root/'Fixture_EN_WithCitations.docx').paragraphs)
        zh=' '.join(p.text for p in Document(self.root/'Fixture_ZH_WithCitations.docx').paragraphs)
        self.assertIn('synthetic fixture',en);self.assertIn('人工合成',zh)

    def test_markdown_source_and_stale_build(self):
        text='The synthetic fixture voltage was 2 V.'
        (self.root/'paragraph.md').write_text(text)
        write(self.root/'content.json',dict(blocks=[dict(id='p1',type='para',files={'en':'paragraph.md'},zh='人工合成测试数据的电压为 2 V。')]))
        self.build();self.assertEqual(self.check().returncode,0)
        (self.root/'paragraph.md').write_text(text+' This is a changed draft.')
        result=self.check();self.assertNotEqual(result.returncode,0);self.assertIn('stale',result.stdout)

    def test_missing_document(self):
        self.build();(self.root/'Fixture_ZH_WithCitations.docx').unlink()
        result=self.check();self.assertNotEqual(result.returncode,0);self.assertIn('document missing',result.stdout)

    def test_release_rejects_unbound_quantity_and_unperformed_reviews(self):
        self.build();result=self.check('--release')
        self.assertNotEqual(result.returncode,0)
        self.assertIn('without bound evidence',result.stdout)
        self.assertIn('review_record',result.stdout)

    def test_table_cell_swap(self):
        write(self.root/'content.json',dict(blocks=[dict(id='p1',type='para',en='Synthetic fixture results are in {tab:t}.',zh='人工合成测试结果见 {tab:t}。')]))
        write(self.root/'floats.json',dict(figures={},boxes={},tables={'t':dict(title_en='Synthetic fixture',title_zh='人工合成测试',columns_en=['A','B'],columns_zh=['A','B'],rows_en=[['2 V','3 V']],rows_zh=[['3 V','2 V']],note_en='',note_zh='')}))
        self.build();self.assertNotEqual(self.check().returncode,0)

    def test_new_project_never_overwrites_local_files(self):
        sentinel=self.root/'local.txt';sentinel.write_text('local edit')
        result=run(ROOT/'skills/sci-manuscript-build/scripts/new_project.py',self.root,'--slug','Other')
        self.assertNotEqual(result.returncode,0);self.assertEqual(sentinel.read_text(),'local edit')

    def test_release_positive_control_and_stale_review(self):
        raw=self.root/'data'/'fixture.json';write(raw,{'voltage':2})
        rec=dict(source_kind='raw_data',value='2',unit='V',metric='fixture voltage',entity='synthetic fixture',result_method='synthetic',source_role='original',access_status='raw-files',verification='verified',provenance=dict(artifacts=[dict(path='data/fixture.json',sha256=sha256(raw))],locator='voltage'))
        write(self.root/'evidence'/'fixture.json',{'v':rec})
        content=json.loads((self.root/'content.json').read_text())
        claim={k:rec[k] for k in ('value','unit','metric','entity','result_method')}
        claim.update(evidence_key='v',text={l:content['blocks'][0][l] for l in ('en','zh')})
        content['blocks'][0]['claims']=[claim];write(self.root/'content.json',content)
        self.build()
        for kind in ('scientific_review','visual_review'):
            result=run(self.root/'record_review.py','--kind',kind,'--reviewed-by','Synthetic test attestation (not actual review)','--notes','Automated fixture for review-record freshness, not a scientific or visual certification',cwd=self.root)
            self.assertEqual(result.returncode,0,result.stderr)
        result=self.check('--release');self.assertEqual(result.returncode,0,result.stdout+result.stderr)
        content['blocks'][0]['en']+=' Synthetic fixture revision.'
        write(self.root/'content.json',content);self.build()
        result=self.check('--release');self.assertNotEqual(result.returncode,0);self.assertIn('review is stale',result.stdout)

    def test_configured_heading_fonts(self):
        self.build()
        from docx import Document
        from docx.oxml.ns import qn
        doc=Document(self.root/'Fixture_ZH_WithCitations.docx')
        cfg=json.loads((self.root/'manuscript.json').read_text())
        self.assertEqual(doc.styles['Title']._element.rPr.rFonts.get(qn('w:eastAsia')),cfg['fonts']['east_asian_heading'])

    def test_default_style_and_opt_in_profiles(self):
        cfg=json.loads((self.root/'manuscript.json').read_text());self.assertEqual(cfg['style']['profile'],'precise')
        default=json.loads((self.root/'style_limits.json').read_text())
        self.assertEqual(default['sentence']['median_min'],16);self.assertEqual(default['sentence']['median_max'],20)
        cfg['style']['profile']='generic';write(self.root/'manuscript.json',cfg);self.build()
        self.assertEqual(self.check().returncode,0)
        self.assertEqual(cfg['writing_context']['perspectives'],['domain researcher','scientific editor'])

    def test_cross_domain_case_bank_and_private_history_exclusion(self):
        folder=ROOT/'skills/sci-writing-orchestrator/case-studies'
        cases=json.loads((folder/'case-manifest.json').read_text())
        self.assertEqual(len(cases),16)
        for case in cases:
            self.assertFalse(case['executed_full_study']);self.assertTrue((folder/case['file']).is_file())
        private_markers=('/Users/','/home/')
        for file in (ROOT/'skills').rglob('*'):
            if file.suffix in ('.md','.json','.py','.csv','.yaml') and '__pycache__' not in file.parts:
                for marker in private_markers:self.assertNotIn(marker,file.read_text(),str(file))

class FigureAndInstallTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()

    def test_missing_final_fails(self):
        folder=self.root/'figure';folder.mkdir()
        write(folder/'figure_spec.json',dict(grid=dict(rows=1,cols=1),slots=[dict(panel='a',kind='data',row=0,col=0,file='panel.png')]))
        result=run(ROOT/'skills/sci-figure-design/scripts/figure_kit.py','verify',folder)
        self.assertNotEqual(result.returncode,0)

    def test_exact_pixels_and_color_changes(self):
        from PIL import Image
        folder=self.root/'figure';folder.mkdir()
        Image.new('RGB',(50,50),(200,50,20)).save(folder/'panel.png')
        canvas=Image.new('RGB',(80,80),'white');canvas.paste(Image.open(folder/'panel.png'),(10,10));canvas.save(folder/'final.png')
        write(folder/'figure_spec.json',dict(grid=dict(rows=1,cols=1),slots=[dict(panel='a',kind='data',row=0,col=0,file='panel.png')]))
        write(folder/'final_placements.json',{'a':dict(box=[10,10,50,50],source_sha256=sha256(folder/'panel.png'))})
        tool=ROOT/'skills/sci-figure-design/scripts/figure_kit.py'
        self.assertEqual(run(tool,'verify',folder).returncode,0)
        canvas.putpixel((15,15),(201,50,20));canvas.save(folder/'final.png')
        self.assertNotEqual(run(tool,'verify',folder).returncode,0)

    def test_install_preserves_previous_skill(self):
        spec=importlib.util.spec_from_file_location('suite_install',ROOT/'tools/install.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        source=self.root/'source';source.mkdir();(source/'SKILL.md').write_text('new version')
        destination=self.root/'installed'/'example';destination.mkdir(parents=True);(destination/'local.txt').write_text('my customization')
        module.place(source,destination,'copy')
        backups=list((destination.parent/'.academic-writing-backups').glob('*/example/local.txt'))
        self.assertEqual(len(backups),1);self.assertEqual(backups[0].read_text(),'my customization')
        self.assertEqual((destination/'SKILL.md').read_text(),'new version')

    def test_missing_composition_slot_fails(self):
        folder=self.root/'figure';folder.mkdir()
        write(folder/'figure_spec.json',dict(grid=dict(rows=1,cols=1),slots=[dict(panel='a',kind='schematic',row=0,col=0,scene='synthetic fixture')]))
        result=run(ROOT/'skills/sci-figure-design/scripts/figure_kit.py','compose',folder)
        self.assertNotEqual(result.returncode,0)
        self.assertFalse((folder/'final.png').exists())

    def test_model_layout_input_never_contains_data_panel(self):
        from PIL import Image
        folder=self.root/'figure';folder.mkdir()
        Image.new('RGB',(80,80),(255,0,255)).save(folder/'data.png')
        write(folder/'figure_spec.json',dict(grid=dict(rows=1,cols=1),slots=[dict(panel='a',kind='data',row=0,col=0,file='data.png')]))
        result=run(ROOT/'skills/sci-figure-design/scripts/figure_kit.py','layout',folder)
        self.assertEqual(result.returncode,0,result.stderr)
        for name,contains_data in [('layout_blocks.png',False),('layout_preview.png',True)]:
            pixels=Image.open(folder/name).convert('RGB').tobytes()
            self.assertEqual(bytes((255,0,255)) in pixels,contains_data)
        self.assertEqual(run(ROOT/'skills/sci-figure-design/scripts/figure_kit.py','prompts',folder).returncode,0)
        prompts=(folder/'prompt_NBP.md').read_text()
        self.assertIn('Never upload data panels',prompts)
        self.assertNotIn('every file in `panels/`',prompts)

    def test_render_failure_is_nonzero_and_stale_pdf_removed(self):
        from unittest.mock import patch
        spec=importlib.util.spec_from_file_location('render_check',ROOT/'skills/sci-manuscript-build/scripts/render_check.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        out=self.root/'qa';out.mkdir();(out/'missing.pdf').write_bytes(b'stale')
        with patch.object(module,'via_soffice',return_value=False),patch.object(module,'configure_fonts'):
            self.assertEqual(module.main([str(self.root/'missing.docx'),'--out',str(out)]),1)
        self.assertFalse((out/'missing.pdf').exists())

if __name__=='__main__':unittest.main()
