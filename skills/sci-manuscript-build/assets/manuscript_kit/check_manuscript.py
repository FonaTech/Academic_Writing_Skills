#!/usr/bin/env python3
"""Consistency, evidence and style checks for a manuscript built with this kit.

Usage
  python3 check_manuscript.py            full report; exit code 1 when a hard check fails
  python3 check_manuscript.py --style    style report only (no documents needed)

Hard checks (FAIL)
  unknown or unused references; figures or tables never cited; equations cited before display;
  abbreviations used before definition; leftover markup in the documents; citation numbers left
  in the no-citation variant; missing language text; citation keys, float references or numbers
  that differ between languages; missing data panels
Soft checks (WARN)
  sentence length profile; overlong sentences and paragraphs; banned phrases, capped words and
  flagged openers from style_limits.json (the file prose_lint.py also reads; change entries per
  project under style.overrides in manuscript.json); sentence openers that repeat; references that
  are not Crossref-verified; body numbers that appear in no evidence record; data panels that may
  have been redrawn by an image model
"""
from __future__ import annotations

import json
import re
import statistics
import sys
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import manuscript_core as rc  # noqa: E402
import style_limits as sl  # noqa: E402
from project_source import load_project, fingerprint, audit_blocks
from evidence_tools import numeric_tokens, quantities, load_evidence, verify_record, audit_bindings, sha256

CFG = json.loads((HERE / 'manuscript.json').read_text(encoding='utf-8'))
STYLE = CFG.get('style', {})
FAIL, WARN = [], []
RELEASE = '--release' in sys.argv or CFG.get('validation', {}).get('mode') == 'release'
PRIMARY = CFG['output'].get('languages', ['en'])[0]


def load_limits():
    """style_limits.json with this project's style.overrides (older median_words_* keys still work)."""
    overrides = dict(STYLE.get('overrides', {}))
    for old, new in (('median_words_min', 'median_min'), ('median_words_max', 'median_max'),
                     ('max_sentence_words', 'max_words')):
        if old in STYLE:
            overrides.setdefault(new, STYLE[old])
    if 'limits' in STYLE:
        WARN.append('style: manuscript.json "limits" is no longer read; move changes to style.overrides '
                    '(keyed by the labels in style_limits.json)')
    return sl.load(sl.profile_path(HERE,STYLE.get('profile','precise')), overrides)


LIMITS = load_limits()


def plain(text):
    text = rc.CITE_RE.sub('', text or '')
    text = rc.FLOAT_RE.sub(lambda m: f'{m.group(1)} X', text)
    text = re.sub(r'\$[^$]*\$', 'X', text)
    text = text.replace('~', '').replace('**', '')
    return re.sub(r'\s+', ' ', text).strip()


def sentences(text):
    protected = re.sub(r'\b(et al|e\.g|i\.e|Fig|Figs|Eq|Eqs|Ref|Refs|ref|vs|approx|ca|cf|no)\.',
                       lambda m: m.group(0).replace('.', '§'), text)
    parts = re.split(r'(?<=[.!?])\s+(?=[A-Z0-9(])', protected)
    return [p.replace('§', '.') for p in parts if len(p.split()) > 2]


def load(name):
    import importlib
    return importlib.import_module(name)


def paragraphs(content, floats, lang='en'):
    for block in content.BLOCKS:
        if block['type'] == 'para':
            yield block['id'], block.get(lang, ''), block
        elif block['type'] == 'box':
            for i, it in enumerate(block['items']):
                if it['type'] == 'para':
                    yield f"{block['id']}.{i}", it.get(lang, ''), it


# --------------------------------------------------------------------- style
def check_style(content, floats):
    if PRIMARY != 'en':
        return
    lengths, long_s, long_p, openers = [], [], [], Counter()
    body, capped, advisory = [], {}, {}
    target = LIMITS['sentence']
    max_words = target['max_words']
    for bid, en, _ in paragraphs(content, floats):
        p = plain(en)
        body.append(p)
        words = len(p.split())
        if words > target['paragraph_max_words']:
            long_p.append((bid, words))
        for entry, m in sl.banned_hits(LIMITS, p):
            WARN.append(f'style: banned "{m.group(0)}" ({entry["label"]}) in {bid}; {entry["fix"]}')
        for group, store in (('caps', capped), ('advisory', advisory)):
            for entry, m in sl.find(LIMITS.get(group, []), p):
                store.setdefault(entry['label'], [entry, []])[1].append(bid)
        for s in sentences(p):
            n = len(s.split())
            lengths.append(n)
            if n > max_words:
                long_s.append((bid, n, s))
            first = ' '.join(s.split()[:2]).rstrip(',')
            openers[first] += 1
            hit = sl.opener_hits(LIMITS, s)
            if hit:
                WARN.append(f'style: {hit[1]} in {bid}; {hit[2]}')
    captions = [(f'fig:{k}', f.get('caption_en', '')) for k, f in floats.FIGURES.items()]
    captions += [(f'tab:{k}', ' '.join([t.get('title_en', ''), t.get('note_en', '')])) for k, t in floats.TABLES.items()]
    for where, text in captions:
        for entry, m in sl.banned_hits(LIMITS, plain(text)):
            WARN.append(f'style: banned "{m.group(0)}" ({entry["label"]}) in {where}; {entry["fix"]}')
    if not lengths:
        print('No English body text yet.')
        return
    med = statistics.median(lengths)
    print(f'English sentences {len(lengths)}; median {med:.0f} words; mean {statistics.mean(lengths):.1f}; '
          f'max {max(lengths)}; >{max_words} words: {len(long_s)}; <=8 words: {sum(n <= 8 for n in lengths)}')
    lo, hi = target['median_min'], target['median_max']
    if not lo <= med <= hi:
        WARN.append(f'style: median sentence length {med:.0f} outside {lo}-{hi} words')
    for bid, n, s in long_s:
        WARN.append(f'long sentence ({n} words) in {bid}: {s[:110]}…')
    for bid, n in long_p:
        WARN.append(f'long paragraph ({n} words) in {bid}; split at the change of topic')
    for first, n in openers.most_common(5):
        if n >= 6 and first.split()[0] not in ('The', 'A', 'An', 'In', 'For', 'This', 'These'):
            WARN.append(f'style: {n} sentences open with "{first}"')
    words = len(' '.join(body).split())
    print(f'English body words (paragraphs and boxes): {words}')
    for kind, store in (('cap', capped), ('check', advisory)):
        for label, (entry, where) in store.items():
            cap = sl.scaled_cap(entry, words, LIMITS.get('reference_words', 10000))
            if len(where) > cap:
                blocks = ', '.join(dict.fromkeys(where))
                WARN.append(f'style: {label} ×{len(where)} ({kind} {cap:g} for {words} words) in {blocks}; {entry["fix"]}')


def _defined(abbr, text):
    """True if text defines abbr as '(ABBR)' / '(term; ABBR)' or, in captions and notes, 'ABBR, term'."""
    a = re.escape(abbr)
    return bool(re.search(rf'\([^()]*(?<![\w-]){a}s?(?![\w-])[^()]*\)', text) or
                re.search(rf'(?:^|[;.]\s|\*\*\s?)\b{a}s?, [a-z]', text))


def check_abbreviations(content, floats):
    """First use in running text needs a definition; so does first use inside each caption and table note."""
    order = [(b['id'], b.get('en', '')) for b in content.BLOCKS if b['type'] == 'para']
    for b in content.BLOCKS:
        if b['type'] == 'box':
            order += [(b['id'], it.get('en', '')) for it in b['items'] if it['type'] == 'para']
    seen_text = [(bid, plain(t)) for bid, t in order]
    abbrs = STYLE.get('abbreviations', [])
    for abbr in abbrs:
        pattern = rf'(?<![\w-]){re.escape(abbr)}s?(?![\w-])'
        first = next(((bid, t) for bid, t in seen_text if re.search(pattern, t)), None)
        if first and not _defined(abbr, first[1]):
            FAIL.append(f'abbreviation {abbr} first used in {first[0]} without a definition')
    # captions and table notes stand alone: each one defines what it uses
    standalone = [(f'fig:{k}', f.get('caption_en', '')) for k, f in floats.FIGURES.items()]
    for k, t in floats.TABLES.items():
        cells = ' '.join(str(c) for row in t.get('rows_en', []) for c in row)
        standalone.append((f'tab:{k}', ' '.join([t.get('title_en', ''), ' '.join(t.get('columns_en', [])), cells,
                                                 t.get('note_en', '')])))
    for where, text in standalone:
        text = plain(text)
        for abbr in abbrs:
            if re.search(rf'(?<![\w-]){re.escape(abbr)}s?(?![\w-])', text) and not _defined(abbr, text):
                WARN.append(f'{where}: {abbr} is used but not defined in the caption or table note')


# --------------------------------------------------------------------- build and documents
def check_build(content):
    path = HERE / 'build_report.json'
    if not path.exists():
        FAIL.append('build_report.json missing: run build_manuscript.py first')
        return None
    report = json.loads(path.read_text(encoding='utf-8'))
    if report.get('source_fingerprint') != fingerprint(HERE, CFG):
        FAIL.append('build is stale: source/evidence/artifacts changed; rebuild before checking')
    if RELEASE and report.get('languages') != CFG['output'].get('languages', ['en']):
        FAIL.append('release requires every requested language output')
    documents = report.get('documents', {})
    if not documents:
        FAIL.append('build report contains no documents')
    for name, digest in documents.items():
        file = HERE/name
        if not file.is_file():
            FAIL.append(f'document missing: {name}')
        elif sha256(file) != digest:
            FAIL.append(f'document changed since build: {name}')
    refs = json.loads((HERE / CFG['paths']['references']).read_text(encoding='utf-8')) \
        if (HERE / CFG['paths']['references']).exists() else {}
    if report.get('unused_references'):
        FAIL.append('unused references: ' + ', '.join(report['unused_references']))
    for key in report.get('references', {}):
        rec = refs.get(key, {})
        if not rec.get('doi'):
            WARN.append(f'reference {key} has no DOI')
        if rec.get('verified') not in ('crossref', 'datacite', 'manual-checked'):
            (FAIL if RELEASE else WARN).append(f'reference {key}: identity not verified by a supported registry or documented manual check')
        if rec.get('post_publication_updates') and (rec.get('integrity_status') != 'reviewed' or not rec.get('integrity_notes')):
            (FAIL if RELEASE else WARN).append(f'reference {key}: publication updates require source inspection and integrity_notes')
    displayed = set()
    for block in content.BLOCKS:
        if block['type'] == 'eq':
            displayed.add(block['key'])
        texts = [block.get(PRIMARY, '')] if block['type'] == 'para' else []
        if block['type'] == 'box':
            for it in block['items']:
                if it['type'] == 'eq':
                    displayed.add(it['key'])
                else:
                    texts.append(it.get(PRIMARY, ''))
        for t in texts:
            for kind, key, _ in rc.FLOAT_RE.findall(t):
                if kind == 'eq' and key not in displayed:
                    FAIL.append(f'equation {key} cited in {block["id"]} before it is displayed')
    print(f"references {len(report.get('references', {}))}, figures {len(report.get('figures', {}))}, "
          f"tables {len(report.get('tables', {}))}, equations {len(report.get('equations', {}))}")
    return report


def check_documents(floats):
    try:
        from docx import Document
    except ImportError:
        WARN.append('python-docx missing; document checks skipped')
        return
    leftovers = re.compile(r'\{(fig|tab|eq):|\[@|\$|(?<![0-9])~|\\[a-z]+\{?|\*\*')
    slug = CFG['project']['slug']
    for path in sorted(HERE.glob(f'{slug}_*.docx')):
        doc = Document(path)
        texts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for r in t.rows:
                texts += [c.text for c in r.cells]
        bad = [t for t in texts if leftovers.search(t)]
        if bad:
            FAIL.append(f'{path.name}: {len(bad)} paragraphs with unresolved markup, e.g. {bad[0][:90]}')
        if 'NoCitation' in path.name:
            nums = re.findall(r'\[\d[\d,–-]*\]', ' '.join(texts))
            if nums:
                FAIL.append(f'{path.name}: citation numbers left: {nums[:5]}')
            if any(re.search(r'\b[Rr]efs?\.\s*(?=[,.;)]|$)', t) for t in texts):
                FAIL.append(f'{path.name}: "ref." left without its citation; rephrase the sentence')
        if 'References' not in path.name:
            pending = sum(1 for t in texts if re.match(r'^\[(Fig\.|图)\s*\d+: artwork pending\]$', t.strip()))
            shown = len(doc.inline_shapes)
            if shown + pending != len(floats.FIGURES):
                FAIL.append(f'{path.name}: {shown} images and {pending} placeholders for {len(floats.FIGURES)} figures')
            elif pending:
                (FAIL if RELEASE else WARN).append(f'{path.name}: {pending} figure(s) still without artwork')


# --------------------------------------------------------------------- languages and numbers
def numbers(text):
    return sorted(numeric_tokens(text))


def compare_pair(where, a, b, primary, lang):
    if sorted(set(rc.parse_cite_keys(a))) != sorted(set(rc.parse_cite_keys(b))):
        FAIL.append(f'{lang}: citation keys differ in {where}')
    if sorted(quantities(a)) != sorted(quantities(b)):
        FAIL.append(f'{lang}: value/unit expressions differ in {where}')
    na, nb = numbers(a), numbers(b)
    if na != nb:
        FAIL.append(f'{lang}: numbers differ in {where}: only {primary} {sorted(set(na) - set(nb))[:6]}, '
                    f'only {lang} {sorted(set(nb) - set(na))[:6]}')


def check_languages(content, floats):
    langs = CFG['output'].get('languages', ['en'])
    if len(langs) < 2:
        return
    primary, others = langs[0], langs[1:]
    for lang in others:
        missing = []
        for bid, text, obj in paragraphs(content, floats, primary):
            other = obj.get(lang, '')
            if not other:
                missing.append(bid)
                continue
            if sorted(set(rc.parse_cite_keys(text))) != sorted(set(rc.parse_cite_keys(other))):
                FAIL.append(f'{lang}: citation keys differ in {bid}')
            fa = sorted(re.sub(r'\|.*\}', '}', m.group(0)) for m in rc.FLOAT_RE.finditer(text))
            fb = sorted(re.sub(r'\|.*\}', '}', m.group(0)) for m in rc.FLOAT_RE.finditer(other))
            if fa != fb:
                FAIL.append(f'{lang}: figure/table/equation references differ in {bid}')
            if sorted(quantities(text)) != sorted(quantities(other)):
                FAIL.append(f'{lang}: value/unit expressions differ in {bid}')
            na, nb = numbers(text), numbers(other)
            if na != nb:
                FAIL.append(f'{lang}: numbers differ in {bid}: only {primary} {sorted(set(na) - set(nb))[:6]}, '
                            f'only {lang} {sorted(set(nb) - set(na))[:6]}')
        for block in content.BLOCKS:
            if block['type'] == 'heading' and not block.get(lang):
                missing.append(block['id'])
        for key, fig in floats.FIGURES.items():
            if not fig.get(f'caption_{lang}'):
                missing.append(f'fig:{key}')
        for key, fig in floats.FIGURES.items():
            a, b = fig.get(f'caption_{primary}', ''), fig.get(f'caption_{lang}', '')
            if b:
                compare_pair(f'fig:{key} caption', a, b, primary, lang)
        for key, tab in floats.TABLES.items():
            ra, rb = tab.get(f'rows_{primary}', []), tab.get(f'rows_{lang}', [])
            if len(rb) != len(ra):
                FAIL.append(f'{lang}: table {key} has a different number of rows')
                continue
            for i, (row_a, row_b) in enumerate(zip(ra, rb)):
                if len(row_a) != len(row_b):
                    FAIL.append(f'{lang}: table {key} row {i} has different column count')
                for ci, (ca, cb) in enumerate(zip(row_a, row_b)):
                    compare_pair(f'tab:{key} row {i+1} cell {ci+1}', str(ca), str(cb), primary, lang)
            compare_pair(f'tab:{key} note', tab.get(f'note_{primary}', ''), tab.get(f'note_{lang}', ''), primary, lang)
        if missing:
            FAIL.append(f'{lang}: {len(missing)} blocks without text: {", ".join(missing[:8])}')


def check_evidence(content, floats):
    langs = CFG['output'].get('languages', ['en'])
    try:
        records = load_evidence(HERE/CFG['paths'].get('evidence', 'evidence'))
    except (ValueError, KeyError) as exc:
        FAIL.append(str(exc))
        return
    blocks = list(audit_blocks(content, floats, langs))
    used = {claim.get('evidence_key') for _,block,_ in blocks for claim in block.get('claims', [])}
    visiting, visited = set(), set()
    def visit(key):
        if key in visiting:
            FAIL.append(f'evidence {key}: cyclic derived inputs')
            return
        if key in visited or key not in records:
            return
        visiting.add(key)
        for source in records[key].get('provenance', {}).get('inputs', []):
            used.add(source)
            visit(source)
        visiting.remove(key)
        visited.add(key)
    for key in list(used):
        visit(key)
    for key, record in records.items():
        for error in verify_record(record, HERE, HERE/CFG['paths'].get('source_audit', 'source_audit'), RELEASE):
            (FAIL if RELEASE and key in used else WARN).append(f'evidence {key}: {error}')
        for source in record.get('provenance', {}).get('inputs', []):
            if source not in records or source == key:
                FAIL.append(f'evidence {key}: invalid derived input {source}')
    errors, notices, checked = audit_bindings(blocks, records, langs, RELEASE)
    (FAIL if RELEASE else WARN).extend(errors)
    WARN.extend(notices)
    print(f'evidence records {len(records)}; explicit claim bindings checked {checked}')
    if RELEASE:
        if not blocks:
            FAIL.append('release has no manuscript body')
        for location, block, texts in blocks:
            if block.get('status') == 'draft' or any('[to verify]' in t for t in texts.values()):
                FAIL.append(f'{location}: unresolved draft/scaffold text')


def check_figures(floats):
    from PIL import Image
    for key, fig in floats.FIGURES.items():
        folder = HERE/CFG['paths']['figures']/fig['folder']
        if (folder/'final_nbp.png').exists():
            (FAIL if RELEASE else WARN).append(f'{key}: model composite is not an accepted data-figure final')
        spec = folder/'figure_spec.json'
        if not spec.exists():
            if RELEASE and not any((folder/n).is_file() for n in ('final.png','final/figure.png')):
                FAIL.append(f'{key}: final artwork missing')
            continue
        data = json.loads(spec.read_text(encoding='utf-8'))
        slots = [slot for slot in data.get('slots', []) if slot.get('kind') == 'data']
        if not slots:
            continue
        final, placements = folder/'final.png', folder/'final_placements.json'
        if not final.is_file() or not placements.is_file():
            (FAIL if RELEASE else WARN).append(f'{key}: data figure needs a code-composed final and placements')
            continue
        placed = json.loads(placements.read_text(encoding='utf-8'))
        if RELEASE and set(placed) != {s['panel'] for s in data.get('slots', [])}:
            FAIL.append(f'{key}: final composition has missing or extra panels')
        big = Image.open(final).convert('RGB')
        for slot in slots:
            item = placed.get(slot['panel'])
            source = folder/slot['file']
            if not isinstance(item, dict) or not source.is_file() or item.get('source_sha256') != sha256(source):
                FAIL.append(f'{key}/{slot["panel"]}: source panel or composition provenance missing/changed')
                continue
            x,y,w,h = item['box']
            expected = Image.open(source).convert('RGB').resize((w,h), Image.Resampling.LANCZOS)
            if big.crop((x,y,x+w,y+h)).tobytes() != expected.tobytes():
                FAIL.append(f'{key}/{slot["panel"]}: data pixels differ from deterministic composition')
        if RELEASE and not (folder/'panels/panel_data.json').is_file():
            FAIL.append(f'{key}: panel_data.json missing; pixel integrity does not prove data correctness')


def check_review():
    path = HERE/'review_record.json'
    if not path.is_file():
        FAIL.append('release requires review_record.json after scientific and visual inspection')
        return
    record = json.loads(path.read_text(encoding='utf-8'))
    if record.get('source_fingerprint') != fingerprint(HERE, CFG):
        FAIL.append('review is stale; inspect the revised manuscript again')
    report_path = HERE/'build_report.json'
    if not report_path.is_file():
        FAIL.append('review requires a current build_report.json')
        return
    report = json.loads(report_path.read_text(encoding='utf-8'))
    if record.get('documents') != report.get('documents'):
        FAIL.append('review does not cover the current built documents')
    for kind in ('scientific_review','visual_review'):
        item = record.get(kind, {})
        if item.get('status') != 'passed' or not item.get('reviewed_by') or not item.get('notes'):
            FAIL.append(f'{kind}: record reviewer, scope and remaining limitations after actual inspection')


def main(argv):
    content, floats = load_project(HERE, CFG)
    check_style(content, floats)
    if '--style' not in argv:
        check_abbreviations(content, floats)
        check_build(content)
        check_documents(floats)
        check_languages(content, floats)
        check_evidence(content, floats)
        check_figures(floats)
        if RELEASE:
            check_review()
    print()
    for w in WARN:
        print('WARN ', w)
    for f in FAIL:
        print('FAIL ', f)
    print(f'\n{len(FAIL)} failures, {len(WARN)} warnings')
    return 1 if FAIL else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
