#!/usr/bin/env python3
"""Build a bilingual, citation-numbered manuscript into Word documents.

Usage
  python3 build_manuscript.py              all languages, all variants
  python3 build_manuscript.py --en-only    authoritative language only
  python3 build_manuscript.py --check      build, then run check_manuscript.py

Inputs (this folder)
  manuscript.json   configuration (titles, languages, fonts, paths, style limits)
  content.py        BLOCKS: ordered headings, paragraphs, display equations, boxes
  floats.py         FIGURES, TABLES, BOXES
  references.json   verified reference database (crossref_verify.py --emit-references)

Outputs
  <slug>_EN_WithCitations.docx, <slug>_EN_NoCitation.docx, <slug>_ZH_WithCitations.docx,
  <slug>_References.docx, <slug>_Equations.json, build_report.json

Pass 1 expands the content in render order, places each figure and table after the
paragraph that first cites it, numbers floats by first mention, equations by display
order and references by first citation. Pass 2 renders each language variant.
"""
from __future__ import annotations

import importlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import manuscript_core as rc  # noqa: E402
from project_source import load_project, fingerprint
from evidence_tools import sha256

CFG = json.loads((HERE / 'manuscript.json').read_text(encoding='utf-8'))
FIG_DIR = HERE / CFG['paths']['figures']
PRIMARY = CFG['output'].get('languages', ['en'])[0]
_OMML_NS = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
_OMML_CACHE: dict = {}
LABELS = {
    'en': {'fig': 'Fig.', 'tab': 'Table', 'refs': 'References', 'pending': '[Layout preview: schematic panels pending]'},
    'zh': {'fig': '图', 'tab': '表', 'refs': '参考文献', 'pending': '[布局预览：示意面板待生成]'},
}


def load_references():
    path = HERE / CFG['paths']['references']
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding='utf-8'))


REFERENCES = load_references()
for _lang in CFG['output'].get('languages', ['en']):
    LABELS.setdefault(_lang, CFG.get('labels', {}).get(_lang, LABELS['en']))


def pandoc_binary() -> Path:
    configured = CFG['paths'].get('pandoc')
    if configured:
        return Path(configured)
    found = shutil.which('pandoc')
    if found:
        return Path(found)
    try:
        import pypandoc
        return Path(pypandoc.get_pandoc_path())
    except Exception:
        pass
    # bundled binaries left by earlier projects (pypandoc-binary inside a local venv)
    for candidate in sorted(HERE.parent.glob('**/pypandoc/files/pandoc'))[:1] + \
            sorted(Path.home().glob('.local/lib/python3*/site-packages/pypandoc/files/pandoc'))[:1]:
        if candidate.exists():
            return candidate
    raise FileNotFoundError('pandoc not found: install pandoc (brew install pandoc) or pypandoc-binary, '
                            'or set paths.pandoc in manuscript.json')


# --------------------------------------------------------------------------
# Low-level docx helpers
# --------------------------------------------------------------------------
def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            el = borders.find(qn(f'w:{edge}'))
            if el is None:
                el = OxmlElement(f'w:{edge}')
                borders.append(el)
            for key in ('val', 'sz', 'space', 'color'):
                if key in kwargs[edge]:
                    el.set(qn(f'w:{key}'), str(kwargs[edge][key]))


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def three_line_table(doc, title, columns, rows, note=None, widths=None):
    """Journal-style three-line table with a bold title line and a table note."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    rc.add_rich_text(p, title)
    for run in p.runs:
        run.font.size = Pt(9.5)
    if p.runs:
        p.runs[0].bold = True
    table = doc.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = widths or [6.7 / len(columns)] * len(columns)
    for i, col in enumerate(columns):
        cell = table.rows[0].cells[i]
        cell.text = ''
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc.add_rich_text(para, col)
        for run in para.runs:
            run.bold = True
            run.font.size = Pt(8)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, 'E4EEF2')
        set_cell_border(cell, top={'val': 'single', 'sz': '12', 'color': '000000'},
                        bottom={'val': 'single', 'sz': '6', 'color': '000000'})
        cell.width = Inches(widths[i])
    repeat_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        last = ridx == len(rows) - 1
        for i, value in enumerate(row):
            cells[i].text = ''
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            rc.add_rich_text(para, str(value))
            for run in para.runs:
                run.font.size = Pt(7.5)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cells[i].width = Inches(widths[i])
            set_cell_border(cells[i], bottom={'val': 'single', 'sz': '12' if last else '2',
                                              'color': '000000' if last else 'BFBFBF'})
    if note:
        q = doc.add_paragraph()
        q.paragraph_format.space_before = Pt(3)
        rc.add_rich_text(q, note)
        for run in q.runs:
            run.font.size = Pt(8.5)
    return table


def equation_omml(key, tex):
    if key in _OMML_CACHE:
        return deepcopy(_OMML_CACHE[key])
    with tempfile.TemporaryDirectory(prefix='kit_eq_') as td:
        out = Path(td) / 'eq.docx'
        subprocess.run([str(pandoc_binary()), '-f', 'markdown', '-t', 'docx', '--standalone', '-o', str(out), '-'],
                       input=f'$$\\displaystyle {tex}$$\n'.encode('utf-8'), check=True,
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        with zipfile.ZipFile(out) as z:
            xml = etree.fromstring(z.read('word/document.xml'))
    elem = xml.find('.//m:oMath', namespaces=_OMML_NS)
    if elem is None:
        raise ValueError(f'pandoc produced no OMML for equation {key}: {tex}')
    _OMML_CACHE[key] = deepcopy(elem)
    return deepcopy(elem)


def add_equation(doc, key, tex, number, zh=False, indent=0.0):
    """Centred OMML equation with a right-aligned number in a borderless two-cell table."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Inches(6.0 - indent)
    right.width = Inches(0.7)
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(equation_omml(key, tex))
    q = right.paragraphs[0]
    q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    q.add_run(f'（{number}）' if zh else f'({number})')
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def figure_image(fig):
    folder = FIG_DIR / fig['folder']
    spec = folder/'figure_spec.json'
    has_data = spec.exists() and any(s.get('kind') == 'data' for s in json.loads(spec.read_text()).get('slots', []))
    names = ('final.png','layout_preview.png') if has_data else ('final/figure.png','final.png','layout_preview.png')
    return next((folder/n for n in names if (folder/n).is_file()), None)


def figure_width(fig):
    """Print width from the figure spec (mm), capped at the text width in manuscript.json."""
    cap = CFG['output'].get('figure_width_in', 6.6)
    spec = FIG_DIR / fig['folder'] / 'figure_spec.json'
    if spec.exists():
        try:
            mm = json.loads(spec.read_text(encoding='utf-8')).get('print_width_mm')
            if mm:
                return min(cap, mm / 25.4)
        except ValueError:
            pass
    return cap


def add_figure(doc, fig, number, caption, lang):
    image = figure_image(fig)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    if image is None:
        rc.add_placeholder(doc, f'[{LABELS[lang]["fig"]} {number}: artwork pending]')
    else:
        p.add_run().add_picture(str(image), width=Inches(figure_width(fig)))
        if image.name.startswith('layout_preview'):
            flag = doc.add_paragraph()
            flag.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = flag.add_run(LABELS[lang]['pending'])
            r.italic = True
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string('B4501E')
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cap.paragraph_format.space_after = Pt(10)
    sep = '' if lang == 'zh' else ' '
    head = cap.add_run(f"{LABELS[lang]['fig']}{sep}{number} | ")
    head.bold = True
    rc.add_rich_text(cap, caption)
    for run in cap.runs:
        run.font.size = Pt(9)
    return cap


# --------------------------------------------------------------------------
# Pass 1
# --------------------------------------------------------------------------
def expand(content, floats):
    """Render-ordered items; each float is released after the first complete paragraph citing it."""
    figs, tabs = floats.FIGURES, floats.TABLES
    placed, queue, items = set(), [], []

    def flush():
        items.extend(queue)
        queue.clear()

    blocks = content.BLOCKS
    for idx, block in enumerate(blocks):
        kind = block['type']
        if kind in ('heading', 'box'):
            flush()
        items.append((kind, block))
        texts = [block[PRIMARY]] if kind == 'para' else (
            [it[PRIMARY] for it in block['items'] if it['type'] == 'para'] if kind == 'box' else [])
        for text in texts:
            for fkind, key, _ in rc.FLOAT_RE.findall(text):
                tag = f'{fkind}:{key}'
                if fkind in ('fig', 'tab') and tag not in placed:
                    pool = figs if fkind == 'fig' else tabs
                    if key not in pool:
                        raise KeyError(f'{fkind} {key} cited in {block["id"]} is not defined in floats.py')
                    placed.add(tag)
                    queue.append((fkind, key))
        nxt = blocks[idx + 1]['type'] if idx + 1 < len(blocks) else None
        if kind in ('para', 'box') and queue and nxt != 'eq':
            flush()
    flush()
    for key in figs:
        if f'fig:{key}' not in placed:
            raise KeyError(f'figure {key} is defined but never cited')
    for key in tabs:
        if f'tab:{key}' not in placed:
            raise KeyError(f'table {key} is defined but never cited')
    return items


def number_sections(content):
    """Number headings and boxes in document order.

    content.CHAPTER (default 1) is the number of the first level-1 heading; each further level-1
    heading adds one (4, 5, 6 ...). Level-2 and level-3 headings number within their chapter
    (4.1, 4.1.1). Headings that already start with a number keep it; unnumbered=True skips one;
    CHAPTER = None switches automatic numbering off.
    """
    first = getattr(content, 'CHAPTER', 1)
    counters = [0, 0, 0]
    sec_no, box_no, n_box = {}, {}, 0
    for block in content.BLOCKS:
        if block['type'] == 'heading':
            level = block.get('level', 1)
            leading = re.match(r'^(\d+(?:\.\d+)*)\s', block[PRIMARY])
            if leading:
                sec_no[block['id']] = leading.group(1)
                continue
            if first is None or level > 3 or block.get('unnumbered'):
                sec_no[block['id']] = None   # the heading exists but has no number; {sec:} to it fails clearly
                continue
            counters[level - 1] += 1
            for i in range(level, 3):
                counters[i] = 0
            chapter = first + max(counters[0], 1) - 1
            sec_no[block['id']] = '.'.join([str(chapter)] + [str(c) for c in counters[1:level]])
        elif block['type'] == 'box':
            n_box += 1
            box_no[block['id']] = n_box
            if block.get('name'):
                box_no[block['name']] = n_box
    return sec_no, box_no


def heading_text(block, lang, sec_no):
    text = block.get(lang, '')
    number = sec_no.get(block['id'])
    return f'{number} {text}' if number and not re.match(r'^\d', text) else text


def number_items(items):
    fig_no, tab_no, eq_no = {}, {}, {}
    for kind, obj in items:
        if kind == 'fig':
            fig_no[obj] = len(fig_no) + 1
        elif kind == 'tab':
            tab_no[obj] = len(tab_no) + 1
        elif kind == 'eq':
            eq_no[obj['key']] = len(eq_no) + 1
        elif kind == 'box':
            for it in obj['items']:
                if it['type'] == 'eq':
                    eq_no[it['key']] = len(eq_no) + 1
    return fig_no, tab_no, eq_no


def text_stream(items, floats, lang='en'):
    """Every citable string in render order: body, boxes, captions, table titles, cells and notes."""
    for kind, obj in items:
        if kind in ('para', 'heading'):
            yield obj['id'], obj.get(lang, '')
        elif kind == 'box':
            yield obj['id'], obj.get(f'title_{lang}', '')
            for it in obj['items']:
                if it['type'] == 'para':
                    yield obj['id'], it.get(lang, '')
        elif kind == 'fig':
            yield f'fig:{obj}', floats.FIGURES[obj].get(f'caption_{lang}', '')
        elif kind == 'tab':
            tab = floats.TABLES[obj]
            yield f'tab:{obj}', tab.get(f'title_{lang}', '')
            for row in tab.get(f'rows_{lang}', []):
                for cell in row:
                    yield f'tab:{obj}', str(cell)
            yield f'tab:{obj}', tab.get(f'note_{lang}', '')


def number_references(items, floats):
    order = {}
    for where, text in text_stream(items, floats, PRIMARY):
        for key in rc.parse_cite_keys(text):
            if key not in REFERENCES:
                raise KeyError(f'unknown citation key @{key} in {where}; add it to references.json')
            order.setdefault(key, len(order) + 1)
    return order


# --------------------------------------------------------------------------
# Pass 2: rendering
# --------------------------------------------------------------------------
def make_resolver(maps, lang, citations):
    ref_order, fig_no, tab_no, eq_no, sec_no, box_no = maps
    style = CFG['output'].get('citation_marker', 'bracket')

    def resolve(text):
        text = rc.resolve_floats(text, fig_no, tab_no, eq_no, lang == 'zh', sec_no, box_no)
        return rc.resolve_citations(text, ref_order, citations, style)
    return resolve


def format_reference(n, ref, style='nature'):
    """Nature-style by default; acs and ieee differ in ordering and emphasis."""
    authors = (ref.get('authors') or '').rstrip('.')
    title = (ref.get('title') or '').rstrip('.')
    journal = ref.get('journal') or ''
    vol, pages, year = ref.get('volume'), ref.get('pages') or ref.get('article_number'), ref.get('year')
    if style == 'ieee':
        out = f'[{n}] {authors}, “{title},” *{journal}*'
        out += f', vol. {vol}' if vol else ''
        out += f', {pages}' if pages else ''
        out += f', {year}.'
    elif style == 'acs':
        out = f'({n}) {authors}. {title}. *{journal}* **{year}**'
        out += f', *{vol}*' if vol else ''
        out += f', {pages}' if pages else ''
        out += '.'
    else:
        out = f'{n}. {authors}. {title}. *{journal}*'
        if vol and pages:
            out += f' **{vol}**, {pages}'
        elif vol:
            out += f' **{vol}**'
        elif pages:
            out += f', {pages}'
        out += f' ({year}).'
    if ref.get('doi'):
        out += f' https://doi.org/{ref["doi"]}'
    return out


def add_references(doc, ref_order, lang):
    doc.add_heading(LABELS[lang]['refs'], level=1)
    style = CFG['output'].get('reference_style', 'nature')
    for key, n in sorted(ref_order.items(), key=lambda kv: kv[1]):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        rc.add_rich_text(p, format_reference(n, REFERENCES[key], style))
        for run in p.runs:
            run.font.size = Pt(10)


def render_box(doc, obj, lang, resolve, eq_no, box_no):
    head = doc.add_paragraph()
    head.paragraph_format.space_before = Pt(10)
    head.paragraph_format.keep_with_next = True
    title = obj.get(f'title_{lang}', '')
    number = box_no.get(obj['id'])
    if number and not re.match(r'^(Box|专栏)\s*\d', title):
        title = f"{'专栏' if lang == 'zh' else 'Box'} {number} | {title}"
    run = head.add_run(resolve(title))
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(rc.TEAL)
    for it in obj['items']:
        if it['type'] == 'para':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            rc.add_rich_text(p, resolve(it.get(lang, '')))
            for r in p.runs:
                r.font.size = Pt(11)
        elif it['type'] == 'eq':
            add_equation(doc, it['key'], it['latex'], eq_no[it['key']], lang == 'zh', indent=0.3)
    end = doc.add_paragraph()
    end.paragraph_format.space_after = Pt(8)
    r = end.add_run('— ' * 3)
    r.font.color.rgb = RGBColor.from_string('9AAAB2')


def render_document(path, content, floats, items, maps, lang='en', citations=True):
    ref_order, fig_no, tab_no, eq_no, sec_no, box_no = maps
    resolve = make_resolver(maps, lang, citations)
    doc = Document()
    rc.set_doc_defaults(doc, CFG.get('fonts', {}))
    fonts = CFG.get('fonts', {})
    normal = doc.styles['Normal']
    normal.font.name = fonts.get('latin', 'Times New Roman')
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), fonts.get('east_asian_body', '宋体'))
    normal.font.size = Pt(fonts.get('body_pt', 12))
    normal.paragraph_format.line_spacing = fonts.get('line_spacing', 1.5)
    project = CFG['project']
    rc.add_title(doc, project.get(f'title_{lang}', ''), project.get(f'subtitle_{lang}') or None)
    for kind, obj in items:
        if kind == 'heading':
            doc.add_heading(resolve(heading_text(obj, lang, sec_no)), level=obj.get('level', 1))
        elif kind == 'para':
            rc.add_body(doc, resolve(obj.get(lang, '')))
        elif kind == 'eq':
            add_equation(doc, obj['key'], obj['latex'], eq_no[obj['key']], lang == 'zh')
        elif kind == 'box':
            render_box(doc, obj, lang, resolve, eq_no, box_no)
        elif kind == 'fig':
            fig = floats.FIGURES[obj]
            add_figure(doc, fig, fig_no[obj], resolve(fig.get(f'caption_{lang}', '')), lang)
        elif kind == 'tab':
            tab = floats.TABLES[obj]
            label = f"{LABELS[lang]['tab']}{'' if lang == 'zh' else ' '}{tab_no[obj]} | "
            rows = [[resolve(str(c)) for c in row] for row in tab.get(f'rows_{lang}', [])]
            three_line_table(doc, label + resolve(tab.get(f'title_{lang}', '')),
                             tab.get(f'columns_{lang}', []), rows,
                             resolve(tab.get(f'note_{lang}', '')) or None, tab.get('widths'))
    if citations and ref_order:
        add_references(doc, ref_order, lang)
    doc.save(path)
    print('  wrote', path.name)


def build_reference_list(path, ref_order):
    doc = Document()
    rc.set_doc_defaults(doc, CFG.get('fonts', {}))
    add_references(doc, ref_order, 'en')
    doc.save(path)
    print('  wrote', path.name)


def main(argv):
    content, floats = load_project(HERE, CFG)
    items = expand(content, floats)
    fig_no, tab_no, eq_no = number_items(items)
    ref_order = number_references(items, floats)
    sec_no, box_no = number_sections(content)
    maps = (ref_order, fig_no, tab_no, eq_no, sec_no, box_no)
    slug = CFG['project']['slug']
    langs = CFG['output'].get('languages', ['en'])
    if '--en-only' in argv or '--primary-only' in argv:
        langs = langs[:1]
    for lang in langs:
        tag = lang.upper()
        render_document(HERE / f'{slug}_{tag}_WithCitations.docx', content, floats, items, maps, lang, True)
        if lang == langs[0] and CFG['output'].get('no_citation_variant', True):
            render_document(HERE / f'{slug}_{tag}_NoCitation.docx', content, floats, items, maps, lang, False)
    if CFG['output'].get('references_only', True) and ref_order:
        build_reference_list(HERE / f'{slug}_References.docx', ref_order)
    equations = {}
    for kind, obj in items:
        eqs = [obj] if kind == 'eq' else [it for it in obj['items'] if it['type'] == 'eq'] if kind == 'box' else []
        for eq in eqs:
            equations[str(eq_no[eq['key']])] = {'key': eq['key'], 'latex': eq['latex']}
    (HERE / f'{slug}_Equations.json').write_text(json.dumps(equations, ensure_ascii=False, indent=2), encoding='utf-8')
    report = {'references': ref_order, 'figures': fig_no, 'tables': tab_no, 'equations': eq_no,
              'unused_references': sorted(set(REFERENCES) - set(ref_order)), 'languages': langs,
              'source_fingerprint': fingerprint(HERE, CFG),
              'documents': {p.name: sha256(p) for p in HERE.glob(f'{slug}_*.docx')}}
    (HERE / 'build_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Built: {len(ref_order)} references, {len(fig_no)} figures, {len(tab_no)} tables, '
          f'{len(eq_no)} equations; languages {", ".join(langs)}.')
    if report['unused_references']:
        print('Unused reference keys:', ', '.join(report['unused_references']))
    if '--check' in argv:
        return subprocess.call([sys.executable, str(HERE / 'check_manuscript.py')] + (['--release'] if '--release' in argv else []))
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
