#!/usr/bin/env python3
"""Assign library ids and folders, then build the library tables described in references/library-layout.md.

Usage
  build_library_tables.py assign search/candidates.csv --library <Topic>_Library/
  fetch_open_access.py <Topic>_Library/library_plan.csv --out <Topic>_Library/ --email you@uni.edu
  build_library_tables.py tables --library <Topic>_Library/ [--topic Topic]

assign  keeps rows whose decision is include (include, yes, y, 1, 纳入), gives each a stable id
        (R### review or perspective, P### primary, X### preprint; an optional 'kind' column
        R/P/X overrides the guess from type, venue and title) and a folder
        01_综述与观点/<category>/, 02_原始研究/<category>/ or 03_预印本/<category>/. Ids already in
        library_plan.csv are kept, matched by DOI. Writes library_plan.csv and selection_manifest.json.
tables  joins the plan with library_status.csv and download_attempts.jsonl and writes
        全部文献.csv, 已下载文献.csv, 未下载文献.csv (UTF-8 BOM), <Topic>_文献库.xlsx, 文献目录.md,
        未下载文献_入口与摘要.md, PDF完整性报告.csv, library.json, and README_draft.md when no
        README.md exists yet. The Chinese notes come from the note_zh column; write them from the
        title and abstract, and say so in the README.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import fitz

sys.path.insert(0, str(Path(__file__).resolve().parent))
from fetch_open_access import validate  # noqa: E402

TOP = {'R': '01_综述与观点', 'P': '02_原始研究', 'X': '03_预印本'}
KIND_NAME = {'R': 'review / perspective', 'P': 'primary research', 'X': 'preprint'}
INCLUDE = ('include', 'yes', 'y', '1', '纳入')
FIELDS = ['id', 'kind', 'category', 'title', 'year', 'venue', 'doi', 'cited_by', 'note_zh', 'abstract',
          'status', 'file', 'entry_points']


def read_csv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as fh:
        return list(csv.DictReader(fh))


def write_csv(path, rows, fields):
    with Path(path).open('w', encoding='utf-8-sig', newline='') as fh:
        wr = csv.DictWriter(fh, fieldnames=fields, extrasaction='ignore')
        wr.writeheader()
        wr.writerows(rows)


def guess_kind(row):
    kind = (row.get('kind') or '').strip().upper()[:1]
    if kind in TOP:
        return kind, 'kind column'
    text = f"{row.get('type', '')} {row.get('venue', '')}".lower()
    if re.search(r'preprint|posted-content|arxiv|biorxiv|chemrxiv|medrxiv|ssrn|research square', text):
        return 'X', 'type or venue'
    if 'review' in text or re.search(r'\b(review|perspective|survey|tutorial|roadmap|progress in|advances in)\b',
                                     row.get('title', '').lower()):
        return 'R', 'type or title (check)'
    return 'P', 'default'


def safe(name):
    return re.sub(r'[\\/:*?"<>|\[\]]+', '_', name).strip() or '未分类'


def assign(args):
    lib = Path(args.library)
    lib.mkdir(parents=True, exist_ok=True)
    plan_path = lib / 'library_plan.csv'
    old = {r['doi'].lower(): r['id'] for r in read_csv(plan_path)} if plan_path.exists() else {}
    used = Counter()   # highest number already given per prefix, so new ids never reuse an old one
    for i in old.values():
        used[i[0]] = max(used[i[0]], int(i[1:]))
    rows, manifest = [], []
    for r in read_csv(args.candidates):
        keep = (r.get('decision') or '').strip().lower() in INCLUDE
        manifest.append(dict(search_id=r.get('id'), doi=r.get('doi'), title=r.get('title'),
                             decision=r.get('decision', ''), category=r.get('category', '')))
        if not keep or not r.get('doi'):
            continue
        kind, basis = guess_kind(r)
        doi = r['doi'].replace('https://doi.org/', '').strip().lower()
        rid = old.get(doi)
        if not rid:
            used[kind] += 1
            rid = f'{kind}{used[kind]:03d}'
        category = safe(r.get('category') or '未分类')
        rows.append(dict(r, id=rid, search_id=r.get('id'), doi=doi, kind=rid[0], kind_basis=basis,
                         category=category, folder=f'{TOP[rid[0]]}/{category}'))
        manifest[-1]['library_id'] = rid
    fields = ['id', 'search_id', 'kind', 'kind_basis', 'category', 'folder', 'title', 'year', 'venue', 'type', 'doi',
              'cited_by', 'is_oa', 'oa_url', 'openalex', 'note_zh', 'abstract']
    write_csv(plan_path, sorted(rows, key=lambda r: r['id']), fields)
    (lib / 'selection_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'{len(rows)} included: ' + ', '.join(f'{k} {v}' for k, v in sorted(Counter(r["kind"] for r in rows).items())))
    print(f'Check kind_basis "(check)" rows, then: fetch_open_access.py {plan_path} --out {lib}')


def local(lib, file):
    for p in (Path(file), lib / file):
        if file and p.exists():
            return p
    return None


def tables(args):
    lib = Path(args.library)
    topic = args.topic or lib.name.replace('_Library', '')
    plan = read_csv(lib / 'library_plan.csv')
    status = {r['id']: r for r in read_csv(lib / 'library_status.csv')} if (lib / 'library_status.csv').exists() else {}
    tried = defaultdict(list)
    if (lib / 'download_attempts.jsonl').exists():
        for line in (lib / 'download_attempts.jsonl').read_text(encoding='utf-8').splitlines():
            a = json.loads(line)
            if a.get('kind') == 'download' and a['url'] not in tried[a['id']]:
                tried[a['id']].append(a['url'])
    rows, integrity = [], []
    for r in plan:
        s = status.get(r['id'], {})
        pdf = local(lib, s.get('file', '')) if s.get('status') in ('downloaded', 'present') else None
        state = s.get('status', 'not attempted')
        if s and not pdf:
            state = 'file missing' if state in ('downloaded', 'present') else 'not retrieved'
        r = dict(r, status=state, file='')
        if pdf:
            try:
                r['file'] = str(pdf.resolve().relative_to(lib.resolve()))
            except ValueError:
                r['file'] = str(pdf)
            ok, detail = validate(pdf, r)
            with fitz.open(pdf) as doc:
                pages = doc.page_count
            integrity.append(dict(id=r['id'], file=r['file'], pages=pages, check='pass' if ok else 'FAIL', detail=detail,
                                  sha256=hashlib.sha256(pdf.read_bytes()).hexdigest()))
        links = [f'https://doi.org/{r["doi"]}'] + [u for u in (r.get('openalex'), r.get('oa_url')) if u]
        r['entry_points'] = ' ; '.join(dict.fromkeys(links + tried.get(r['id'], [])[:4]))
        rows.append(r)
    have = [r for r in rows if r['file']]
    missing = [r for r in rows if not r['file']]
    write_csv(lib / '全部文献.csv', rows, FIELDS)
    write_csv(lib / '已下载文献.csv', have, FIELDS)
    write_csv(lib / '未下载文献.csv', missing, FIELDS)
    write_csv(lib / 'PDF完整性报告.csv', integrity, ['id', 'file', 'pages', 'check', 'detail', 'sha256'])
    (lib / 'library.json').write_text(json.dumps({r['id']: dict(r, attempts=tried.get(r['id'], [])) for r in rows},
                                                 ensure_ascii=False, indent=1), encoding='utf-8')
    workbook(lib / f'{topic}_文献库.xlsx', rows)
    catalogue(lib, rows, missing)
    if not (lib / 'README.md').exists():
        readme(lib, topic, rows, have, integrity)
    print(f'{len(rows)} papers, {len(have)} with a local PDF, {len(missing)} without; tables written to {lib}')


def workbook(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    head = ['ID', 'Type', 'Category', 'Title', 'Year', 'Venue', 'DOI', 'Cited by', '中文说明', 'Abstract', 'Status',
            'Local PDF']
    wb = Workbook()
    ov = wb.active
    ov.title = '概览'
    ov.append(['Type', 'Category', 'Papers', 'With PDF'])
    for (kind, cat), group in _groups(rows).items():
        ov.append([KIND_NAME[kind], cat, len(group), sum(bool(r['file']) for r in group)])
    sheets = [('全部', rows)] + [(f'{k}_{c}', g) for (k, c), g in _groups(rows).items()]
    names = set()
    for name, group in sheets:
        name = safe(name)[:28]
        while name in names:
            name = name[:26] + f'_{len(names)}'
        names.add(name)
        ws = wb.create_sheet(name)
        ws.append(head)
        for r in group:
            ws.append([r['id'], KIND_NAME[r['kind']], r['category'], r['title'], r.get('year'), r.get('venue'),
                       r['doi'], r.get('cited_by'), r.get('note_zh', ''), (r.get('abstract') or '')[:1500],
                       r['status'], r['file']])
            n = ws.max_row
            ws.cell(n, 7).hyperlink, ws.cell(n, 7).style = f'https://doi.org/{r["doi"]}', 'Hyperlink'
            if r['file']:
                ws.cell(n, 12).hyperlink, ws.cell(n, 12).style = r['file'], 'Hyperlink'
        for col, width in zip('ABCDEFGHIJKL', (7, 16, 16, 60, 6, 28, 28, 8, 40, 70, 12, 40)):
            ws.column_dimensions[col].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes, ws.auto_filter.ref = 'A2', ws.dimensions
    wb.save(path)


def _groups(rows):
    """Rows grouped by (kind, category) in folder order: reviews, primary research, preprints."""
    groups = defaultdict(list)
    for r in sorted(rows, key=lambda r: (list(TOP).index(r['kind']), r['category'], r['id'])):
        groups[(r['kind'], r['category'])].append(r)
    return groups


def catalogue(lib, rows, missing):
    out = ['# 文献目录', '']
    for (kind, cat), group in _groups(rows).items():
        out += [f'## {TOP[kind]} / {cat}', '']
        for r in group:
            where = f'[PDF]({r["file"]})' if r['file'] else '未下载'
            note = f' {r["note_zh"]}' if r.get('note_zh') else ''
            out.append(f'- {r["id"]} {r["title"]} ({r.get("year", "")}, {r.get("venue", "")}). '
                       f'https://doi.org/{r["doi"]} · {where}.{note}')
        out.append('')
    (lib / '文献目录.md').write_text('\n'.join(out), encoding='utf-8')
    out = ['# 未下载文献：入口与摘要', '',
           '未下载不等于闭源；可经图书馆、出版社页面或作者主页获取。没有绕过任何付费墙或验证。', '']
    for r in missing:
        out += [f'## {r["id"]} {r["title"]}', '', f'{r.get("year", "")}, {r.get("venue", "")}; status: {r["status"]}',
                '', 'Entry points: ' + r['entry_points'], '', (r.get('abstract') or 'No abstract in the search record.'), '']
    (lib / '未下载文献_入口与摘要.md').write_text('\n'.join(out), encoding='utf-8')


def readme(lib, topic, rows, have, integrity):
    kinds = Counter(r['kind'] for r in rows)
    failed = [i['id'] for i in integrity if i['check'] != 'pass']
    text = [f'# {topic} literature library', '',
            'Scope: <fill in>. Search date: <fill in>. Reading route: <foundational review, then key primary papers, '
            'then extensions>.', '',
            f'Papers: {len(rows)} ({", ".join(f"{KIND_NAME[k]} {n}" for k, n in sorted(kinds.items()))}); '
            f'with a local PDF: {len(have)}; without: {len(rows) - len(have)}.', '',
            f'PDF integrity check failures: {", ".join(failed) if failed else "none"}.', '',
            'Caveats: the Chinese notes are based on title and abstract, not a full quality review. '
            '"Not downloaded" does not mean closed access, and no paywall or verification page was bypassed. '
            'Citation counts are snapshots on the search date.', '']
    (lib / 'README_draft.md').write_text('\n'.join(text), encoding='utf-8')


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    a = sub.add_parser('assign', help='ids and folders for included candidates')
    a.add_argument('candidates')
    a.add_argument('--library', required=True)
    t = sub.add_parser('tables', help='CSV, xlsx and Markdown deliverables')
    t.add_argument('--library', required=True)
    t.add_argument('--topic')
    args = ap.parse_args(argv)
    (assign if args.cmd == 'assign' else tables)(args)
    return 0


if __name__ == '__main__':
    sys.exit(main())
