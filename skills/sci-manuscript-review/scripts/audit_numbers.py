#!/usr/bin/env python3
"""Audit individual claim-to-evidence bindings; never use a global number pool as support."""
from __future__ import annotations
import argparse
import glob
import importlib.util
import json
import sys
from pathlib import Path

KIT = Path(__file__).resolve().parents[2]/'sci-manuscript-build/assets/manuscript_kit'
sys.path.insert(0,str(KIT))
from evidence_tools import iter_records, load_evidence, audit_bindings
from project_source import load_project, audit_blocks

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('sources',nargs='*')
    ap.add_argument('--project')
    ap.add_argument('--evidence',nargs='+')
    ap.add_argument('--release',action='store_true')
    ap.add_argument('--trust-python',action='store_true',help='explicitly execute legacy local content modules')
    ap.add_argument('--json',action='store_true')
    args = ap.parse_args(argv)
    if not args.project and not args.sources:
        ap.error('supply --project or sources')
    records = {}
    if args.evidence:
        for pattern in args.evidence:
            for filename in glob.glob(pattern):
                for key,rec in iter_records(json.loads(Path(filename).read_text(encoding='utf-8'))):
                    if key in records:ap.error('duplicate evidence key '+key)
                    records[key] = rec
    if args.project:
        root = Path(args.project).resolve()
        cfg = json.loads((root/'manuscript.json').read_text(encoding='utf-8'))
        sys.path.insert(0,str(root))
        if not args.evidence:
            records = load_evidence(root/cfg['paths'].get('evidence','evidence'))
        content,floats = load_project(root,cfg)
        langs = cfg['output'].get('languages',['en'])
        blocks = list(audit_blocks(content,floats,langs))
    else:
        langs = ['en']
        blocks = []
        for filename in args.sources:
            path = Path(filename)
            if path.suffix == '.docx':
                from docx import Document
                doc = Document(path)
                blocks += [(f'{path.name}:{i}',{},dict(en=p.text)) for i,p in enumerate(doc.paragraphs) if p.text]
                for ti,table in enumerate(doc.tables):
                    for ri,row in enumerate(table.rows):
                        blocks += [(f'{path.name}:T{ti}:{ri}:{ci}',{},dict(en=c.text)) for ci,c in enumerate(row.cells)]
            elif path.suffix == '.py' and args.trust_python:
                sys.path.insert(0,str(path.resolve().parent))
                spec = importlib.util.spec_from_file_location(path.stem,path)
                mod = importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
                blocks += [(b['id'],b,dict(en=b.get('en',''))) for b in getattr(mod,'BLOCKS',[]) if b.get('type')=='para']
            else:
                ap.error('use declarative --project; legacy Python requires --trust-python')
    errors,notices,checked = audit_bindings(blocks,records,langs,True)
    result = dict(bindings_checked=checked,findings=errors,notices=notices,
                  limitation='Declared consistency only; semantic source/data review is separate.')
    print(json.dumps(result,ensure_ascii=False,indent=2) if args.json else '\n'.join(errors+notices)+f'\n{checked} bindings checked; {len(errors)} errors')
    return 1 if errors or not blocks else 0

if __name__ == '__main__':
    sys.exit(main())
